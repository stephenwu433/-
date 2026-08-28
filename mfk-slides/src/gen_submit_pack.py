#!/usr/bin/env python3
"""
Submit pack: fix CJK font tofu, lock cover lexicon, Part03 rename,
build compressed main PPTX (≤32) + appendix PPTX.
"""
from pathlib import Path
import io
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Inches

OUT = Path("/workspace/mfk-slides")
ART = Path("/opt/cursor/artifacts")
W, H = 3840, 2160
MX = 140
GAP = 28
FY = 2055

CREAM = (245, 242, 235)
INK = (28, 28, 28)
MUTED = (90, 90, 90)
RED = (155, 36, 51)
TEAL = (34, 100, 108)
GOLD = (184, 148, 90)
LINE = (210, 205, 196)
WHITE = (255, 255, 255)
SOFT = (236, 232, 224)
SOFT_T = (232, 240, 241)
SOFT_R = (244, 232, 234)
SOFT_G = (245, 238, 224)

BG_COVER = Path("/home/ubuntu/.cursor/projects/workspace/assets/38a58645-eda5-4f5b-9a7e-0e2a7c23d3d2.png")
BG_DARK = Path("/home/ubuntu/.cursor/projects/workspace/assets/ea789c2c-218c-4b09-ab49-46393f3059c0.png")
if not BG_DARK.exists():
    BG_DARK = Path("/home/ubuntu/.cursor/projects/workspace/assets/49a9cc62-4bc4-4ae4-a911-53783998238c.png")

POS = "私人香气表达的精准权威"
MOTIF = "精准的感性"
SYSTEM = "时刻衣柜"
JOURNEY = "从试香到签名"
PROMISE = "让香气，准确表达复杂的你。"
PROBLEM = "消费者欣赏MFK，却难以确定哪一支真正适合自己"

# CJK-safe fonts only for any Chinese text
_S = "/usr/share/fonts/opentype/noto/NotoSerifCJK-Bold.ttc"
_N = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
_B = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
_I = "/usr/share/fonts/truetype/inter/Inter-Regular.ttf"
_IB = "/usr/share/fonts/truetype/inter/Inter-Bold.ttf"


def fonts():
    return {
        "brand": ImageFont.truetype(_I, 22),  # Latin only
        "eyebrow": ImageFont.truetype(_B, 28, index=2),  # CJK
        "title": ImageFont.truetype(_S, 54, index=2),
        "hero": ImageFont.truetype(_S, 68, index=2),
        "h2": ImageFont.truetype(_S, 34, index=2),
        "h3": ImageFont.truetype(_B, 26, index=2),
        "body": ImageFont.truetype(_N, 24, index=2),
        "small": ImageFont.truetype(_N, 20, index=2),
        "tiny": ImageFont.truetype(_N, 16, index=2),
        "sub": ImageFont.truetype(_N, 26, index=2),
        "label": ImageFont.truetype(_B, 18, index=2),  # CJK — was Inter tofu
        "num": ImageFont.truetype(_IB, 36),
        "row": ImageFont.truetype(_N, 30, index=2),
        "row_b": ImageFont.truetype(_B, 32, index=2),
        "foot": ImageFont.truetype(_I, 18),
    }


F = fonts()


def rr(d, box, fill, r=12):
    d.rounded_rectangle(box, radius=r, fill=fill)


def wrap(d, text, xy, font, fill, max_w, lh):
    x, y = xy
    text = str(text).replace("\r", " ").replace("\n", " ")
    line = ""
    for ch in text:
        test = line + ch
        if d.textlength(test, font=font) <= max_w:
            line = test
        else:
            if line:
                d.text((x, y), line, font=font, fill=fill)
                y += lh
            line = ch
    if line:
        d.text((x, y), line, font=font, fill=fill)
    return y


def header(d, code, title, question):
    d.text((MX, 40), "MAISON FRANCIS KURKDJIAN  PARIS", font=F["brand"], fill=MUTED)
    d.text((W - MX - 200, 40), "LVMH BEAUTY", font=F["brand"], fill=MUTED)
    d.text((MX, 88), code, font=F["label"], fill=RED)
    d.text((MX, 118), title, font=F["title"], fill=INK)
    d.rectangle([MX, 198, MX + 100, 204], fill=GOLD)
    d.text((MX, 220), question, font=F["sub"], fill=MUTED)


def footer(d, src):
    d.line([(MX, FY - 16), (W - MX, FY - 16)], fill=LINE, width=2)
    d.text((MX, FY), src, font=F["tiny"], fill=MUTED)
    d.text((W - MX - 360, FY), "MFK Brand Strategy  |  CONFIDENTIAL", font=F["tiny"], fill=MUTED)


def save(img, name):
    img.save(OUT / name, "PNG", optimize=True)
    img.resize((1920, 1080), Image.Resampling.LANCZOS).save(
        OUT / name.replace(".png", "_1920.png"), "PNG", optimize=True
    )
    img.save(ART / name, "PNG", optimize=True)
    print("saved", name)


# ═══════════════════════════════════════════════════════════
# Cover — ONLY one strategic positioning
# ═══════════════════════════════════════════════════════════
def slide_cover():
    bg = Image.open(BG_COVER).convert("RGB").resize((W, H), Image.Resampling.LANCZOS)
    overlay = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    for i in range(0, 1750):
        a = int(90 * (1 - i / 1750))
        od.line([(i, 0), (i, H)], fill=(245, 242, 235, a))
    img = Image.alpha_composite(bg.convert("RGBA"), overlay).convert("RGB")
    d = ImageDraw.Draw(img)

    d.text((MX, 88), "MAISON FRANCIS KURKDJIAN  PARIS", font=F["brand"], fill=MUTED)
    d.text((W - MX - 240, 88), "LVMH BEAUTY", font=F["brand"], fill=MUTED)

    d.text((MX, 420), "LVMH BEAUTY  ·  未来 1—3 年香氛美学与体验创新方案", font=F["eyebrow"], fill=RED)
    wrap(d, f"MFK品牌全案：{PROMISE}", (MX, 520), F["hero"], INK, 1780, 86)
    wrap(d, "以调香师专业与嗅觉衣橱，帮助消费者确认属于自己的标志性香气。", (MX, 740), F["sub"], INK, 1680, 44)

    # ONLY strategic positioning — no second "品牌方向"
    d.text((MX, 920), "战略定位", font=F["label"], fill=MUTED)
    d.text((MX + 220, 912), POS, font=F["row_b"], fill=RED)
    tw = d.textlength(POS, font=F["row_b"])
    d.rectangle([MX + 220, 962, MX + 220 + int(tw), 968], fill=GOLD)

    # Supporting hierarchy as labels — not competing directions
    support = [
        (f"品牌母题  {MOTIF}", MUTED),
        (f"产品系统  {SYSTEM}", MUTED),
        (f"体验机制  {JOURNEY}", MUTED),
    ]
    y = 1040
    for text, col in support:
        d.text((MX + 220, y), text, font=F["body"], fill=col)
        y += 56

    d.text((MX, 1980), "MFK × LVMH BEAUTY  /  BRAND GROWTH SYSTEM", font=F["foot"], fill=MUTED)
    save(img, "mfk_00_cover.png")


def slide_ending():
    bg = Image.open(BG_DARK).convert("RGB").resize((W, H), Image.Resampling.LANCZOS)
    img = bg.copy()
    d = ImageDraw.Draw(img)
    CREAM2 = (245, 242, 235)
    SOFT2 = (200, 188, 178)
    MUTED2 = (160, 148, 140)

    d.text((MX, 88), "MAISON FRANCIS KURKDJIAN  PARIS", font=F["brand"], fill=MUTED2)
    d.text((W - MX - 240, 88), "LVMH BEAUTY", font=F["brand"], fill=MUTED2)

    d.text((MX, 360), "LVMH BEAUTY  ·  未来 1—3 年香氛美学与体验创新方案", font=F["eyebrow"], fill=GOLD)
    wrap(d, f"MFK品牌全案：{PROMISE}", (MX, 450), F["hero"], CREAM2, 1900, 84)
    d.rectangle([MX, 660, MX + 140, 666], fill=GOLD)
    wrap(d, "以调香师专业与嗅觉衣橱，帮助消费者确认属于自己的标志性香气。", (MX, 710), F["sub"], SOFT2, 1800, 44)

    d.text((MX, 880), "战略定位", font=F["label"], fill=MUTED2)
    d.text((MX + 240, 872), POS, font=F["row_b"], fill=GOLD)
    tw = d.textlength(POS, font=F["row_b"])
    d.rectangle([MX + 240, 922, MX + 240 + int(tw), 928], fill=GOLD)

    rows = [
        ("品牌母题", MOTIF),
        ("产品系统", SYSTEM),
        ("体验机制", JOURNEY),
    ]
    y = 980
    for lab, val in rows:
        d.text((MX, y), lab, font=F["label"], fill=MUTED2)
        d.text((MX + 240, y - 2), val, font=F["row"], fill=CREAM2)
        y += 70

    d.text((MX, 1240), "品牌问题", font=F["label"], fill=MUTED2)
    wrap(d, PROBLEM, (MX + 240, 1236), F["body"], SOFT2, 1600, 34)

    d.text((MX, 1400), "下一步", font=F["label"], fill=GOLD)
    d.text((MX + 240, 1396), "内部证据链已通过 → 真人概念验证（非市场独占宣称）", font=F["body"], fill=CREAM2)

    d.text((MX, 1780), "EVIDENCE-LED  ·  HUMAN-VERIFIED  ·  PRIORITIZE VALIDATION", font=F["foot"], fill=MUTED2)
    save(img, "mfk_00_ending.png")


def slide_case_close():
    bg = Image.open(BG_DARK).convert("RGB").resize((W, H), Image.Resampling.LANCZOS)
    img = bg.copy()
    d = ImageDraw.Draw(img)
    CREAM2 = (245, 242, 235)
    SOFT2 = (200, 188, 178)
    MUTED2 = (160, 148, 140)

    d.text((MX, 80), "MAISON FRANCIS KURKDJIAN  PARIS", font=F["brand"], fill=MUTED2)
    d.text((W - MX - 220, 80), "LVMH BEAUTY", font=F["brand"], fill=MUTED2)
    d.text((MX, 200), "00 / CASE CLOSE", font=F["label"], fill=GOLD)
    d.text((MX, 250), "全案收束：评委应记住的唯一层级", font=F["hero"], fill=CREAM2)
    d.rectangle([MX, 350, MX + 120, 356], fill=GOLD)

    beats = [
        ("战略定位", POS),
        ("品牌母题", MOTIF),
        ("产品系统", SYSTEM),
        ("体验机制", JOURNEY),
        ("消费者表达", PROMISE.replace("。", "")),
        ("下一步", "真人概念验证，非市场独占宣称"),
    ]
    bw = (W - 2 * MX - 5 * GAP) // 6
    for i, (a, b) in enumerate(beats):
        x = MX + i * (bw + GAP)
        card = Image.new("RGBA", (bw, 480), (40, 18, 20, 170))
        img.paste(card, (x, 420), card)
        d = ImageDraw.Draw(img)
        d.text((x + 16, 450), a, font=F["label"], fill=GOLD)
        wrap(d, b, (x + 16, 520), F["h3"], CREAM2, bw - 32, 36)

    d = ImageDraw.Draw(img)
    wrap(
        d,
        f"MFK 要成为「{POS}」：用「{MOTIF}」判断，用「{SYSTEM}」编排，用「{JOURNEY}」确认。",
        (MX, 980),
        F["h2"],
        CREAM2,
        W - 2 * MX,
        48,
    )
    d.text((MX, 1140), PROMISE, font=F["hero"], fill=CREAM2)
    wrap(
        d,
        "内部证据链已通过筛选；方向进入真人验证。过闸再扩面——方向成立 ≠ 规模成立。",
        (MX, 1280),
        F["body"],
        SOFT2,
        W - 2 * MX,
        36,
    )
    d.text((MX, 1780), "EVIDENCE-LED  ·  HUMAN-VERIFIED  ·  PRIORITIZE VALIDATION", font=F["foot"], fill=MUTED2)
    save(img, "mfk_00_case_close.png")


def slide_toc():
    img = Image.new("RGB", (W, H), CREAM)
    d = ImageDraw.Draw(img)
    header(d, "00 / CONTENTS", "目录", "主汇报 29 页；技术中间表见附录。")

    parts = [
        ("PART 01", "界定战场与研究问题", "问题 · 目标消费者 · 战场分层"),
        ("PART 02", "市场、消费者与竞品诊断", "信号 · 摩擦 · 竞品 · 时刻衣柜"),
        ("PART 03", "系统筛选与方向收敛", "从411条语料筛选出两个待真人验证的方向"),
        ("PART 04", "进入点决策与美学预判", "场景主切 · 精准克制 · 锁定移交"),
        ("PART 05", "品牌主张与现实编排", "双声部主张 · 时刻衣柜逻辑 · 四线编排"),
        ("PART 06", "90天试点与评估", "最小闭环 · 节奏 · Go/Iterate/Kill"),
    ]
    for i, (code, title, sub) in enumerate(parts):
        y = 320 + i * 240
        rr(d, [MX, y, W - MX, y + 210], WHITE if i % 2 == 0 else SOFT_T, 10)
        d.text((MX + 40, y + 40), code, font=F["label"], fill=RED if i == 2 else TEAL)
        d.text((MX + 40, y + 90), title, font=F["h2"], fill=INK)
        d.text((MX + 40, y + 150), sub, font=F["body"], fill=MUTED)
        if i == 2:
            d.rectangle([MX, y, MX + 12, y + 210], fill=RED)

    footer(d, "PART 03 已更名：系统筛选与方向收敛（非「战略跃迁与证据验证」）")
    save(img, "mfk_00_toc.png")


def slide_p03_opener():
    """Part 3 home — no Feishu URL; screening framing."""
    img = Image.new("RGB", (W, H), CREAM)
    d = ImageDraw.Draw(img)
    header(
        d,
        "03 / SYSTEM SCREENING",
        "PART 03｜系统筛选与方向收敛",
        "从411条语料中筛选出两个待真人验证的品牌方向——不是市场总验证。",
    )

    rr(d, [MX, 300, W - MX, 560], SOFT_R, 12)
    d.text((MX + 36, 340), "本章一句话", font=F["label"], fill=RED)
    wrap(
        d,
        "内部语料证据链支持将「精准的感性」作母题、「从试香到签名」作体验机制，服务定位「私人香气表达的精准权威」，进入真人概念验证；备选方向保留。",
        (MX + 36, 400),
        F["h3"],
        INK,
        W - 2 * MX - 80,
        44,
    )

    cols = [
        ("输入", "语料池（IG / TikTok / 小红书等）", "语料信号"),
        ("过程", "冲突编码 → 候选方向 → 闸门筛选", "方法设定"),
        ("输出", "BRD-02 母题 + BRD-01 体验（待真人验证）", "策略推断"),
    ]
    cw = (W - 2 * MX - 2 * GAP) // 3
    for i, (a, b, c) in enumerate(cols):
        x = MX + i * (cw + GAP)
        rr(d, [x, 620, x + cw, 1100], WHITE, 10)
        d.text((x + 28, 660), a, font=F["label"], fill=TEAL)
        wrap(d, b, (x + 28, 740), F["h3"], INK, cw - 56, 40)
        d.text((x + 28, 980), "证据状态 · " + c, font=F["small"], fill=GOLD)

    rr(d, [MX, 1180, W - MX, 1560], SOFT_G, 10)
    d.text((MX + 28, 1240), "读数边界", font=F["h3"], fill=GOLD)
    wrap(
        d,
        "系统筛选通过 ≠ 市场验证成立。语料条数、标签数、洞察簇、候选方向是不同对象，禁止读成同一转化漏斗。详细标签词典与中间计数表见附录。",
        (MX + 28, 1320),
        F["body"],
        INK,
        W - 2 * MX - 60,
        40,
    )
    d.text((MX + 28, 1480), "禁止用语：命题成立 / 市场已验证 / 竞品无法替代（绝对句）", font=F["body"], fill=RED)

    rr(d, [MX, 1640, W - MX, 1920], SOFT_T, 10)
    d.text((MX + 28, 1700), "本章主汇报只保留", font=F["label"], fill=TEAL)
    d.text(
        (MX + 28, 1760),
        "① 筛选契约  →  ② 机制规则  →  ③ 主推为何是「精准的感性」  →  ④ 筛选结论",
        font=F["h3"],
        fill=INK,
    )
    d.text((MX + 28, 1840), "编码明细（DIR / CHK / 标签词典）移入附录，不在主线展开。", font=F["body"], fill=MUTED)

    footer(d, "语料来源：Instagram、TikTok、小红书等平台；完整数据表见附录。")
    save(img, "mfk_03_v1_contract.png")


def slide_appendix_divider():
    img = Image.new("RGB", (W, H), CREAM)
    d = ImageDraw.Draw(img)
    header(d, "APPENDIX", "附录｜技术中间表与编码明细", "供答辩追问；主汇报可不翻。")

    items = [
        ("A1", "冲突编码与洞察簇明细", "INS 平台分布 / CF 标签"),
        ("A2", "候选方向生成表", "DIR-01~04 字段还原"),
        ("A3", "闸门矩阵与压力测试", "CHK + 五维"),
        ("A4", "方法管道与漏斗计数", "异对象计数说明"),
        ("A5", "资产证明与结构信号", "可兑现资产对照"),
        ("A6", "竞品深潜补充页", "Byredo / Le Labo"),
    ]
    cw = (W - 2 * MX - 2 * GAP) // 3
    for i, (code, title, sub) in enumerate(items):
        x = MX + (i % 3) * (cw + GAP)
        y = 340 + (i // 3) * 700
        rr(d, [x, y, x + cw, y + 620], WHITE, 12)
        d.text((x + 36, y + 80), code, font=F["label"], fill=RED)
        wrap(d, title, (x + 36, y + 180), F["h2"], INK, cw - 72, 48)
        wrap(d, sub, (x + 36, y + 400), F["body"], MUTED, cw - 72, 36)

    footer(d, "附录不改变主线结论；主线裁定仍是：内部证据链通过 → 真人概念验证")
    save(img, "mfk_99_appendix_divider.png")


# ═══════════════════════════════════════════════════════════
# Build PPTX from JPEG slides (MS PowerPoint friendly, small)
# ═══════════════════════════════════════════════════════════
MAIN_ORDER = [
    ("mfk_00_cover.png", "封面"),
    ("mfk_00_strategic_answer.png", "一眼看懂"),
    ("mfk_00_toc.png", "目录"),
    ("mfk_01_core_problem.png", "核心问题"),
    ("mfk_01_target_consumer.png", "目标消费者"),
    ("01_layer_decision.png", "战场分层"),
    ("mfk_02_market_signal.png", "市场信号"),
    ("mfk_02_consumer_friction.png", "消费者摩擦"),
    ("mfk_02_competitor_map.png", "竞品地图"),
    ("mfk_02_mfk_diagnosis.png", "MFK诊断"),
    ("mfk_02_mfk_proposition.png", "时刻衣柜命题"),
    ("mfk_03_00_chapter.png", "Part03章节扉页"),
    ("mfk_03_v1_contract.png", "系统筛选开篇"),
    ("mfk_03_01_mechanism.png", "筛选机制"),
    ("mfk_03_05_why_brd02.png", "为何主推精准的感性"),
    ("mfk_03_v5_q3_validate.png", "筛选结论"),
    ("mfk_04_01_entry_decision.png", "进入点"),
    ("mfk_04_02_aesthetic_forecast.png", "美学预判"),
    ("mfk_04_03_lock_handoff.png", "锁定移交"),
    ("mfk_05_01_proposition.png", "主张锁定"),
    ("mfk_05_wardrobe_logic.png", "时刻衣柜逻辑"),
    ("mfk_05_02_orchestration_map.png", "四线编排"),
    ("mfk_05_03_product_retail.png", "产品零售"),
    ("mfk_05_04_content_relationship.png", "内容关系"),
    ("mfk_06_01_minimum_loop.png", "最小闭环"),
    ("mfk_06_02_ninety_day_plan.png", "90天计划"),
    ("mfk_06_03_evaluate_decide.png", "评估决策"),
    ("mfk_00_case_close.png", "全案收束"),
    ("mfk_00_ending.png", "结尾"),
]

APPENDIX_ORDER = [
    ("mfk_99_appendix_divider.png", "附录扉页"),
    ("mfk_03_02_conflict_encoding.png", "冲突编码"),
    ("mfk_03_03_ins_to_dir.png", "洞察转方向"),
    ("mfk_03_04_gate_matrix.png", "闸门矩阵"),
    ("mfk_03_v2_mechanism.png", "机制图表"),
    ("mfk_03_v3_q1_conflict.png", "冲突计量"),
    ("mfk_03_v4_q2_gates.png", "闸门计量"),
    ("mfk_03_method_pipeline.png", "方法管道"),
    ("mfk_03_selection_funnel.png", "筛选漏斗"),
    ("mfk_03_insight_clusters.png", "洞察簇"),
    ("mfk_03_asset_proof.png", "资产证明"),
    ("mfk_03_structure_signal.png", "结构信号"),
    ("mfk_02_competitor_deep.png", "竞品深潜"),
]


def png_to_jpeg_bytes(path: Path, max_side=1920, quality=82) -> bytes:
    im = Image.open(path).convert("RGB")
    w, h = im.size
    if w > max_side:
        im = im.resize((max_side, int(h * max_side / w)), Image.Resampling.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=quality, optimize=True)
    return buf.getvalue()


def build_pptx(order, out_path: Path, title: str):
    # 16:9 widescreen — use EMU for exact 1920x1080 slide
    prs = Presentation()
    prs.slide_width = Emu(12192000)  # 13.333" ≈ 1920px at 144dpi mapping
    prs.slide_height = Emu(6858000)  # 7.5"

    blank = prs.slide_layouts[6]
    missing = []
    for fname, _ in order:
        src = OUT / fname
        if not src.exists():
            # try without mfk_ prefix variants
            missing.append(fname)
            continue
        slide = prs.slides.add_slide(blank)
        jpeg = png_to_jpeg_bytes(src, max_side=1920, quality=82)
        # write temp
        tmp = OUT / "_tmp_slide.jpg"
        tmp.write_bytes(jpeg)
        slide.shapes.add_picture(str(tmp), 0, 0, width=prs.slide_width, height=prs.slide_height)
    if (OUT / "_tmp_slide.jpg").exists():
        (OUT / "_tmp_slide.jpg").unlink()

    prs.save(str(out_path))
    size_mb = out_path.stat().st_size / 1e6
    print(f"PPTX {out_path.name}: {len(order) - len(missing)} slides, {size_mb:.1f} MB", "missing:", missing)
    return size_mb, missing


def patch_generators_fonts():
    """Patch Inter→Noto for label/eyebrow in sibling generators (CJK tofu fix)."""
    import re

    files = list((OUT / "src").glob("gen_*.py"))
    for path in files:
        if path.name == "gen_submit_pack.py":
            continue
        t = path.read_text()
        orig = t
        # label / eyebrow that use Inter Bold
        t = re.sub(
            r'"label":\s*ImageFont\.truetype\(ib,\s*(\d+)\)',
            r'"label": ImageFont.truetype(b, \1, index=2)',
            t,
        )
        t = re.sub(
            r'"eyebrow":\s*ImageFont\.truetype\(ib,\s*(\d+)\)',
            r'"eyebrow": ImageFont.truetype(b, \1, index=2)',
            t,
        )
        # gen_clarity_fix uses F with ib for label - also fix if present
        if path.name == "gen_clarity_fix.py":
            t = t.replace(
                '"label": ImageFont.truetype(ib, 18),',
                '"label": ImageFont.truetype(b, 18, index=2),',
            )
            t = t.replace(
                '"eyebrow": ImageFont.truetype(ib, 28),',
                '"eyebrow": ImageFont.truetype(b, 28, index=2),',
            )
        if t != orig:
            path.write_text(t)
            print("font-patched", path.name)


if __name__ == "__main__":
    patch_generators_fonts()
    slide_cover()
    slide_toc()
    slide_p03_opener()
    slide_case_close()
    slide_ending()
    slide_appendix_divider()

    main_path = OUT / "MFK_LVMH_品牌全案_主汇报.pptx"
    app_path = OUT / "MFK_LVMH_品牌全案_附录.pptx"
    mb, miss = build_pptx(MAIN_ORDER, main_path, "MFK Main")
    mb2, miss2 = build_pptx(APPENDIX_ORDER, app_path, "MFK Appendix")

    # also copy to artifacts
    import shutil

    shutil.copy(main_path, ART / main_path.name)
    shutil.copy(app_path, ART / app_path.name)
    print("MAIN_PAGES", len(MAIN_ORDER), "MB", round(mb, 1))
    print("APPENDIX_PAGES", len(APPENDIX_ORDER), "MB", round(mb2, 1))
    print("SUBMIT PACK DONE")
