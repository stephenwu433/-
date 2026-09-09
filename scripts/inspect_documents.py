#!/usr/bin/env python3
"""Inspect the source documents that make up the `ban` repository.

The repo currently contains only planning/specification material for the
"梅见 AI 市场决策系统" (Excel workbooks, a PDF, and Word documents). This
utility parses each supported document and prints a short summary so the
Cloud Agent environment can be exercised end to end without any application
source code yet existing.

Usage:
    python scripts/inspect_documents.py            # summarize every document
    python scripts/inspect_documents.py FILE ...   # summarize specific files
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from docx import Document
import pdfplumber

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

REPO_ROOT = Path(__file__).resolve().parent.parent


def _xlsx_raw_row_counts(path: Path) -> dict[str, tuple[int, int]]:
    """Map each sheet name to (row_count, max_column) from the raw XML.

    openpyxl fails on these workbooks (they contain a dataValidation with an
    empty sqref) and read-only mode trusts a stale `A1` <dimension>, so the
    counts are read straight from the zipped sheet XML.
    """
    counts: dict[str, tuple[int, int]] = {}
    with zipfile.ZipFile(path) as zf:
        wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
        rels_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        rid_to_target = {
            rel.get("Id"): rel.get("Target")
            for rel in rels_root.findall(f"{{{_PKG_REL_NS}}}Relationship")
        }
        for sheet in wb_root.findall(f"{{{_MAIN_NS}}}sheets/{{{_MAIN_NS}}}sheet"):
            name = sheet.get("name")
            rid = sheet.get(f"{{{_REL_NS}}}id")
            target = rid_to_target.get(rid, "")
            if target.startswith("/"):
                part = target[1:]
            elif target.startswith("xl/"):
                part = target
            else:
                part = f"xl/{target}"
            try:
                data = zf.read(part).decode("utf-8", "replace")
            except KeyError:
                counts[name] = (0, 0)
                continue
            rows = len(re.findall(r"<row\b", data))
            max_col = 0
            for ref in re.findall(r'<c\s+r="([A-Z]+)\d+"', data):
                col = 0
                for ch in ref:
                    col = col * 26 + (ord(ch) - ord("A") + 1)
                max_col = max(max_col, col)
            counts[name] = (rows, max_col)
    return counts


def summarize_xlsx(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheetnames = list(wb.sheetnames)
    wb.close()
    raw_counts = _xlsx_raw_row_counts(path)
    lines = [f"  workbook opened with openpyxl; {len(sheetnames)} sheet(s):"]
    for name in sheetnames:
        rows, cols = raw_counts.get(name, (0, 0))
        lines.append(f"    - {name!r}: {rows} rows x {cols} cols")
    return "\n".join(lines)


def summarize_docx(path: Path) -> str:
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    total_chars = sum(len(p) for p in paragraphs)
    lines = [
        f"  {len(paragraphs)} non-empty paragraph(s), "
        f"{len(doc.tables)} table(s), {total_chars} chars"
    ]
    if paragraphs:
        first = paragraphs[0][:80]
        lines.append(f"    first line: {first!r}")
    return "\n".join(lines)


def summarize_pdf(path: Path) -> str:
    with pdfplumber.open(path) as pdf:
        pages = pdf.pages
        first_text = (pages[0].extract_text() or "").strip().splitlines()
        head = first_text[0][:80] if first_text else "(no extractable text on page 1)"
        return f"  {len(pages)} page(s)\n    page 1 starts: {head!r}"


HANDLERS = {
    ".xlsx": summarize_xlsx,
    ".docx": summarize_docx,
    ".pdf": summarize_pdf,
}


def discover_documents() -> list[Path]:
    found: list[Path] = []
    for ext in HANDLERS:
        found.extend(sorted(REPO_ROOT.glob(f"*{ext}")))
    return found


def main(argv: list[str]) -> int:
    if argv:
        targets = [Path(a) for a in argv]
    else:
        targets = discover_documents()

    if not targets:
        print("No supported documents found.")
        return 1

    failures = 0
    for path in targets:
        print(f"\n== {path.name} ==")
        handler = HANDLERS.get(path.suffix.lower())
        if handler is None:
            print(f"  (unsupported file type: {path.suffix})")
            failures += 1
            continue
        try:
            print(handler(path))
        except Exception as exc:  # pragma: no cover - surfaced to the user
            print(f"  ERROR: {exc}")
            failures += 1

    print(f"\nInspected {len(targets)} document(s); {failures} failure(s).")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
